import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def enforce_formatting():
    doc_path = 'GROUP 2 CHAPTER 1 TO CHAPTER 4.docx'
    if not os.path.exists(doc_path):
        print(f"Error: {doc_path} not found.")
        return

    doc = docx.Document(doc_path)

    # 1. Update Default Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(4)

    # Count word and paragraph metrics
    total_words = 0

    # 2. Process all Paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        total_words += len(txt.split())

        # Set paragraph line spacing to 1.5
        p.paragraph_format.line_spacing = 1.5

        # Determine if paragraph is a major heading, sub-heading, or body text
        is_chapter = txt.startswith('CHAPTER') or txt.startswith('CHAPTER ')
        is_major_section = False
        if len(txt) > 0 and txt[0].isdigit() and ('.' in txt[:5] or ' ' in txt[:5]):
            if len(txt) < 80:
                is_major_section = True

        # Process each run in paragraph
        if len(p.runs) == 0 and txt:
            # If text exists without explicit runs
            run = p.add_run(txt)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14 if is_chapter else 12)
        else:
            for r in p.runs:
                r.font.name = 'Times New Roman'
                if is_chapter:
                    r.font.size = Pt(14)
                    r.font.bold = True
                elif is_major_section:
                    r.font.size = Pt(12)
                    r.font.bold = True
                else:
                    r.font.size = Pt(12)

    # 3. Process all Tables
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.5
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(12)

    # 4. Estimate Page Count
    # Standard 12pt Times New Roman, 1.5 line spacing, 1-inch margins averages ~300 words per page.
    # Images and tables also contribute to page count.
    num_tables = len(doc.tables)
    num_paragraphs = len(doc.paragraphs)
    estimated_pages = max(1, int(total_words / 280 + num_tables * 0.75 + 10))

    print(f"Formatting Complete!")
    print(f"Total Words: {total_words}")
    print(f"Total Paragraphs: {num_paragraphs}")
    print(f"Total Tables: {num_tables}")
    print(f"Estimated Page Count: ~{estimated_pages} pages (Constraint: <= 90 pages)")

    doc.save(doc_path)
    print(f"Saved updated document to {doc_path}")

if __name__ == "__main__":
    enforce_formatting()
