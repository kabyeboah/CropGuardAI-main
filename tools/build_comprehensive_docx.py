#!/usr/bin/env python3
"""
CropGuard AI — Comprehensive System Documentation Generator (.docx)
Builds an exhaustive, professional technical architecture, ML pipeline,
database schema, and system implementation document.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

OUTPUT_DOCX = "CropGuard_AI_Comprehensive_System_Documentation.docx"

# Color Palette Constants
HEX_PRIMARY = "1B5E20"      # Dark Forest Green
HEX_SECONDARY = "2E7D32"    # Medium Green
HEX_ACCENT = "388E3C"       # Light Green Accent
HEX_DARK_TEXT = "212121"    # Charcoal Body Text
HEX_LIGHT_BG = "F4F6F4"     # Off-white / light green tint
HEX_BORDER = "CCCCCC"       # Subtle border grey
HEX_WARNING_BG = "FFF8E1"   # Amber tint for warnings
HEX_WARNING_BORDER = "FF8F00"
HEX_INFO_BG = "E3F2FD"      # Blue tint for notes
HEX_INFO_BORDER = "1976D2"
HEX_CODE_BG = "F5F5F5"      # Code block grey

COLOR_PRIMARY = RGBColor(0x1B, 0x5E, 0x20)
COLOR_SECONDARY = RGBColor(0x2E, 0x7D, 0x32)
COLOR_ACCENT = RGBColor(0x38, 0x8E, 0x3C)
COLOR_DARK_TEXT = RGBColor(0x21, 0x21, 0x21)
COLOR_MUTED = RGBColor(0x61, 0x61, 0x61)


def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    tc_pr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    """Sets internal padding (in twips) of a table cell."""
    tc_pr = cell._element.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def set_callout_borders(cell, border_color_hex, border_size="36"):
    """Sets a thick left border and removes top, bottom, right borders."""
    tc_pr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)


def set_table_borders(table, border_color_hex=HEX_BORDER):
    """Applies clean subtle borders to a standard table."""
    tbl_pr = table._element.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color_hex}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_PRIMARY}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color_hex}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)


class DocxBuilder:
    def __init__(self):
        self.doc = docx.Document()
        self._setup_page_geometry()
        self._setup_styles()

    def _setup_page_geometry(self):
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)
            section.page_width = Inches(8.5)
            section.page_height = Inches(11.0)

            # Header / Footer
            footer = section.footer
            f_p = footer.paragraphs[0]
            f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            f_run = f_p.add_run("CropGuard AI — Technical Architecture & Implementation Manual | Confidential")
            f_run.font.name = "Calibri"
            f_run.font.size = Pt(8.5)
            f_run.font.color.rgb = COLOR_MUTED

    def _setup_styles(self):
        # Base normal style
        style_normal = self.doc.styles['Normal']
        font_normal = style_normal.font
        font_normal.name = 'Calibri'
        font_normal.size = Pt(10.5)
        font_normal.color.rgb = COLOR_DARK_TEXT
        style_normal.paragraph_format.line_spacing = 1.15
        style_normal.paragraph_format.space_after = Pt(4)

    def add_title_block(self, title, subtitle, metadata_dict):
        """Creates an executive title cover block."""
        p_top = self.doc.add_paragraph()
        p_top.paragraph_format.space_before = Pt(36)
        p_top.paragraph_format.space_after = Pt(12)
        run_badge = p_top.add_run("CROPGUARD AI — PRODUCTION TECHNICAL MANUAL")
        run_badge.font.name = 'Calibri'
        run_badge.font.size = Pt(11)
        run_badge.font.bold = True
        run_badge.font.color.rgb = COLOR_ACCENT

        p_title = self.doc.add_paragraph()
        p_title.paragraph_format.space_after = Pt(8)
        run_title = p_title.add_run(title)
        run_title.font.name = 'Calibri'
        run_title.font.size = Pt(26)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_PRIMARY

        p_sub = self.doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(24)
        run_sub = p_sub.add_run(subtitle)
        run_sub.font.name = 'Calibri'
        run_sub.font.size = Pt(13)
        run_sub.font.italic = True
        run_sub.font.color.rgb = COLOR_MUTED

        # Meta Table
        table = self.doc.add_table(rows=len(metadata_dict), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        set_table_borders(table, "E0E0E0")
        for i, (k, v) in enumerate(metadata_dict.items()):
            row = table.rows[i]
            cell_k = row.cells[0]
            cell_v = row.cells[1]
            cell_k.width = Inches(2.2)
            cell_v.width = Inches(4.3)
            set_cell_margins(cell_k, 60, 60, 80, 80)
            set_cell_margins(cell_v, 60, 60, 80, 80)
            set_cell_background(cell_k, HEX_LIGHT_BG)
            
            p_k = cell_k.paragraphs[0]
            r_k = p_k.add_run(k)
            r_k.font.bold = True
            r_k.font.size = Pt(9.5)
            r_k.font.color.rgb = COLOR_SECONDARY
            
            p_v = cell_v.paragraphs[0]
            r_v = p_v.add_run(v)
            r_v.font.size = Pt(9.5)
            r_v.font.color.rgb = COLOR_DARK_TEXT

        p_div = self.doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(24)
        p_div.paragraph_format.space_after = Pt(24)
        r_div = p_div.add_run("―" * 45)
        r_div.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)

    def add_h1(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_h3(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_ACCENT
        return p

    def add_p(self, text, bold_prefix=None, italic=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK_TEXT
        r_body = p.add_run(text)
        r_body.font.italic = italic
        r_body.font.color.rgb = COLOR_DARK_TEXT
        return p

    def add_bullet(self, text, bold_prefix=None, level=0):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.bold = True
        r_text = p.add_run(text)
        r_text.font.color.rgb = COLOR_DARK_TEXT
        return p

    def add_callout(self, text, title=None, callout_type="info"):
        """Creates a callout box table with custom left border and shading."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)

        bg_hex = HEX_INFO_BG if callout_type == "info" else (
            HEX_WARNING_BG if callout_type == "warning" else HEX_LIGHT_BG
        )
        border_hex = HEX_INFO_BORDER if callout_type == "info" else (
            HEX_WARNING_BORDER if callout_type == "warning" else HEX_PRIMARY
        )

        set_cell_background(cell, bg_hex)
        set_callout_borders(cell, border_hex, "36")
        set_cell_margins(cell, 140, 140, 180, 180)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15

        if title:
            r_title = p.add_run(f"{title}\n")
            r_title.font.bold = True
            r_title.font.size = Pt(10.5)
            r_title.font.color.rgb = RGBColor(0x0D, 0x47, 0xA1) if callout_type == "info" else (
                RGBColor(0xE6, 0x51, 0x00) if callout_type == "warning" else COLOR_PRIMARY
            )

        r_text = p.add_run(text)
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_DARK_TEXT

        # Space after table
        p_after = self.doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(2)
        p_after.paragraph_format.space_after = Pt(2)

    def add_code_block(self, code_text):
        """Creates an indented, shaded monospace code block."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, HEX_CODE_BG)
        set_cell_margins(cell, 120, 120, 160, 160)

        # Subtle thin grey border all around
        tc_pr = cell._element.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
            f'<w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_ACCENT}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
            f'</w:tcBorders>'
        )
        tc_pr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(code_text.strip())
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x2E, 0x34, 0x40)

        p_after = self.doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(2)
        p_after.paragraph_format.space_after = Pt(2)

    def add_styled_table(self, headers, rows_data, col_widths=None):
        """Adds a professionally formatted data table."""
        table = self.doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, "D0D0D0")

        # Header Row
        hdr_cells = table.rows[0].cells
        for col_idx, header in enumerate(headers):
            cell = hdr_cells[col_idx]
            set_cell_background(cell, HEX_PRIMARY)
            set_cell_margins(cell, 120, 120, 140, 140)
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(header)
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Data Rows
        for row_idx, row_values in enumerate(rows_data):
            row_cells = table.rows[row_idx + 1].cells
            bg_hex = HEX_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(row_values):
                cell = row_cells[col_idx]
                set_cell_background(cell, bg_hex)
                set_cell_margins(cell, 100, 100, 120, 120)
                if col_widths and col_idx < len(col_widths):
                    cell.width = col_widths[col_idx]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                # align left for text, center for short tags/numbers
                str_val = str(val)
                if len(str_val) < 12 and ("%" in str_val or str_val.replace('.', '', 1).isdigit() or str_val in ["Yes", "No", "High", "Low", "Moderate", "FAIL", "PASS"]):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str_val)
                r.font.size = Pt(9.0)
                r.font.color.rgb = COLOR_DARK_TEXT
                if str_val in ["FAIL", "High"]:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
                elif str_val in ["PASS", "Yes"]:
                    r.font.bold = True
                    r.font.color.rgb = COLOR_PRIMARY

        p_after = self.doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(4)
        p_after.paragraph_format.space_after = Pt(4)

    def add_image_figure(self, image_path, caption_title, caption_desc=None, width=Inches(6.2)):
        """Embeds an image centered with professional figure caption."""
        if not os.path.exists(image_path):
            self.add_callout(f"Image asset missing: {image_path}", "Figure Missing", "warning")
            return

        p_img = self.doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=width)

        p_cap = self.doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption_title)
        run_cap.font.bold = True
        run_cap.font.size = Pt(9.5)
        run_cap.font.color.rgb = COLOR_SECONDARY

        if caption_desc:
            run_desc = p_cap.add_run(f" ― {caption_desc}")
            run_desc.font.size = Pt(9.0)
            run_desc.font.italic = True
            run_desc.font.color.rgb = COLOR_MUTED

    def add_page_break(self):
        self.doc.add_page_break()

    def save(self, filepath):
        self.doc.save(filepath)
        print(f"Document successfully written to: {filepath}")


def build_cropguard_documentation():
    b = DocxBuilder()

    # ──────────────────────────────────────────────────────────────────────────
    # TITLE COVER & METADATA
    # ──────────────────────────────────────────────────────────────────────────
    meta = {
        "Document Type": "Comprehensive Engineering Architecture & Implementation Manual",
        "Target Platform": "Flutter (Android / iOS) + Supabase Edge Backend + On-Device TFLite",
        "Target Agronomic Context": "Smallholder Agriculture in Ghana & Sub-Saharan West Africa",
        "ML Model Architecture": "MobileNetV2 (128x128x3) Transfer Learning with Softmax Calibration",
        "Backend Architecture": "Supabase PostgreSQL (RLS), Supabase Storage & Deno Edge Functions",
        "Primary Offline Engine": "SQLite (sqflite v18) + Pending Sync Queue + Local Tile Cache",
        "Author / Engineering Team": "CropGuard AI Core Engineering, Mobile & Agronomy Architecture Team",
        "Date of Documentation": "September 2026 (Production Baseline v4.0)"
    }
    b.add_title_block(
        title="CropGuard AI: Comprehensive Technical Architecture, Machine Learning Pipeline & Implementation Manual",
        subtitle="A Technical Blueprint of an Offline-First Edge AI & Agro-Meteorological Surveillance System for Smallholder Crop Disease Management",
        metadata_dict=meta
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 1: EXECUTIVE SUMMARY & AGRO-TECHNICAL CONTEXT
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("1. Executive Summary & Problem Formulation")
    b.add_p(
        "CropGuard AI is a production-grade, offline-first mobile application designed to empower smallholder farmers "
        "and agricultural extension officers in Ghana and broader Sub-Saharan Africa. Crop diseases—including cassava mosaic virus, "
        "maize streak, tomato blights, and cocoa black pod—cause catastrophic yield losses ranging from 30% to 100% when diagnosed late. "
        "In developing agricultural economies, smallholder farmers face three critical structural barriers:",
        bold_prefix="System Context: "
    )
    b.add_bullet("Severe Agricultural Extension Officer Deficit: In Ghana, the extension officer-to-farmer ratio historically exceeds 1:1,500, meaning smallholders cannot receive timely in-person diagnostic support during critical pathogen incubation periods.")
    b.add_bullet("Remote Field Connectivity Blackouts: Farm fields are frequently located beyond the reach of reliable 3G/4G cellular networks, rendering cloud-only computer vision APIs useless for real-time in-field decisions.")
    b.add_bullet("Literacy & Language Barriers: Rural farmers frequently require vernacular voice guidance (e.g., Twi, Ga, Ewe) and simple, iconographic treatment protocols rather than complex academic agronomy manuals.")

    b.add_p(
        "To solve these operational realities, CropGuard AI implements a multi-tiered engineering paradigm: an ultra-fast on-device "
        "Convolutional Neural Network (MobileNetV2 running via TensorFlow Lite) that executes in sub-120 milliseconds without network access; "
        "an offline SQLite database with an automatic FIFO synchronization engine; an agro-meteorological forecasting engine "
        "that anticipates disease outbreaks triggered by rainfall and humidity; interactive privacy-preserved outbreak maps; "
        "and a secure multimodal cloud AI fallback (Gemini 3 Flash) hosted on Supabase Edge Functions.",
        bold_prefix="Core Value Proposition: "
    )

    b.add_callout(
        "CropGuard AI is engineered around strict zero-data-loss and privacy guarantees. All leaf scans, treatment tracking milestones, "
        "and diagnostic history records are committed to local SQLite before any network dispatch is attempted. Farmer GPS coordinates "
        "are coarsened to ~1.1 km to protect farm boundaries and personal property while providing regional outbreak awareness.",
        "Guiding Agronomic & Architectural Philosophy",
        "info"
    )

    b.add_image_figure(
        "cropguard_screens_preview.png",
        "Figure 1.1: CropGuard AI Mobile Interface Preview",
        "Demonstrating Home Dashboard, Leaf Scanning Viewport, Diagnostic Results, and Community Outbreak Surveillance."
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 2: SYSTEM ARCHITECTURE & DESIGN PATTERNS
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("2. System Architecture & Software Design Patterns")
    b.add_p(
        "CropGuard AI is built following the principles of Clean Architecture combined with the Model-View-ViewModel (MVVM) pattern, "
        "Provider for reactive state management, and GetIt for dependency injection. This architectural separation enforces loose "
        "coupling, high cohesion, testability, and clean boundaries between the UI presentation, domain business logic, and low-level data sources."
    )

    b.add_h2("2.1 Layered Architectural Structure")
    b.add_p(
        "The codebase under lib/ is strictly partitioned into three primary architectural rings, each with distinct responsibilities:"
    )

    b.add_bullet("Presentation Layer (lib/presentation/): Contains Flutter widgets, feature screens, and ChangeNotifier ViewModels. The presentation layer observes domain state and handles user interaction. Crucially, presentation classes NEVER import data sources or perform raw SQL/Supabase queries directly. All actions flow through Use Cases.")
    b.add_bullet("Domain Layer (lib/domain/): The pure Dart business core of the application. It contains plain data models (DetectionResult, TreatmentPlan, WeatherForecast, DiseaseRisk), abstract repository interfaces (IAuthRepository, IDetectionRepository, ICommunityRepository, IWeatherRepository, IRiskRepository), and single-responsibility Use Cases (ScanCropUseCase, ScanBatchUseCase, GetWeatherUseCase). It has zero dependencies on Flutter UI or third-party infrastructure SDKs.")
    b.add_bullet("Data Layer (lib/data/): The implementation layer that satisfies domain contracts. It houses concrete repository implementations (DetectionRepositoryImpl, WeatherRepositoryImpl, RiskRepositoryImpl), local databases (DatabaseHelper with SQLite, PendingSyncQueue), on-device ML engines (CropDiseaseClassifier with tflite_flutter), and remote data sources (SupabaseDatabaseService, SupabaseAuthService, SupabaseStorageService, CloudFunctionsService).")

    b.add_h2("2.2 Dependency Injection via GetIt")
    b.add_p(
        "Dependency injection is centralized in lib/core/di/service_locator.dart via the setupServiceLocator() bootstrapper. "
        "Services are registered either as singletons (for long-lived caches and database connections) or factories (for transient use cases). "
        "This allows seamless substitution of real services with mock implementations during unit and widget testing without altering production code."
    )

    b.add_code_block("""
// lib/core/di/service_locator.dart (Architectural Pattern)
final sl = GetIt.instance;

Future<void> setupServiceLocator() async {
  // 1. Core Services & Database Singletons
  sl.registerLazySingleton<DatabaseHelper>(() => DatabaseHelper());
  sl.registerLazySingleton<CropDiseaseClassifier>(() => CropDiseaseClassifier());
  sl.registerLazySingleton<SupabaseDatabaseService>(() => SupabaseDatabaseService(Supabase.instance.client));

  // 2. Concrete Repositories (bound to Domain Interfaces)
  sl.registerLazySingleton<IDetectionRepository>(
    () => DetectionRepositoryImpl(sl<DatabaseHelper>(), sl<SupabaseDatabaseService>()),
  );
  sl.registerLazySingleton<IWeatherRepository>(
    () => WeatherRepositoryImpl(),
  );
  sl.registerLazySingleton<IRiskRepository>(
    () => RiskRepositoryImpl(sl<ICommunityRepository>(), sl<IWeatherRepository>()),
  );

  // 3. Encapsulated Use Cases
  sl.registerFactory<ScanCropUseCase>(
    () => ScanCropUseCase(sl<IClassifierRepository>(), sl<IDetectionRepository>(), sl<StreakManager>()),
  );
}
""")

    b.add_h2("2.3 Error Handling via Functional Result Pattern")
    b.add_p(
        "Rather than allowing unhandled runtime exceptions to bubble up and crash the UI, CropGuard AI employs a functional "
        "Result<T> pattern defined in lib/core/utils/result.dart. Repositories and Use Cases return either Result.success(data) "
        "or Result.error(failure). Failures are strongly typed domain objects (lib/core/error/failures.dart):"
    )
    b.add_bullet("MLFailure: Captures tensor mismatch, out-of-memory, or corrupt image issues during model inference.")
    b.add_bullet("DatabaseFailure: Captures SQLite lock contention, schema constraint violations, or local disk corruption.")
    b.add_bullet("NetworkFailure: Represents HTTP timeouts, offline status, or failed DNS resolution.")
    b.add_bullet("ServerFailure: Represents upstream Supabase HTTP 5xx responses or Edge Function timeouts.")

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 3: ON-DEVICE AI MODEL: TRAINING, ARCHITECTURE & MOBILE INTEGRATION
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("3. On-Device AI Model: Architecture, Training & Mobile Integration")
    b.add_p(
        "At the core of CropGuard AI's edge capabilities is an on-device Convolutional Neural Network (CNN) trained specifically "
        "to identify 51 distinct plant disease and pest manifestations across tropical and global staples. The model is packaged "
        "as assets/cropguard_plant_disease.tflite and executes natively on the mobile device's CPU/NPU without incurring cellular data costs."
    )

    b.add_h2("3.1 Dataset Aggregation & African Agronomic Curation")
    b.add_p(
        "Standard public plant pathology datasets (such as original PlantVillage) suffer from acute geographic bias: they feature leaves "
        "plucked in temperate climates and photographed on uniform grey backgrounds. When deployed on Ghanaian farms, such models suffer "
        "catastrophic domain shift due to harsh tropical sunlight, red laterite soil backgrounds, weed occlusion, and localized pathogen variants. "
        "To combat this, CropGuard AI's training data was aggregated from multiple diverse field sources:"
    )

    b.add_bullet("Ghana Crop Pest & Disease Dataset: Real field photographs of cassava, cocoa, maize, and tomato collected from Eastern, Ashanti, and Northern farming communities.", bold_prefix="Ghana MoFA / PPRSD Sources: ")
    b.add_bullet("Uganda Cassava Dataset: Field imagery of Cassava Brown Streak Disease (CBSD) and Cassava Mosaic Disease (CMD) under equatorial field illumination.", bold_prefix="Makerere AI Lab: ")
    b.add_bullet("In-Field Maize Pathogen Corpus: Extensive photos of Northern Corn Leaf Blight, Common Rust, and Grey Leaf Spot with authentic background clutter.", bold_prefix="CGIAR / IITA Maize Hub: ")
    b.add_bullet("Rice Disease Field Dataset: Brown Spot, Rice Blast, and Leaf Scald photographed directly in irrigated paddy fields.", bold_prefix="AfricaRice Corpus: ")
    b.add_bullet("Specialty Crop Repositories: High-resolution field leaves for Cashew Anthracnose, Mango Malformation/Anthracnose, Sugarcane Rust, Banana Sigatoka, and Groundnut Leaf Spot.", bold_prefix="Regional Corpora: ")

    b.add_h2("3.2 Preprocessing, Data Augmentation & Model Architecture")
    b.add_p(
        "The model architecture is built upon MobileNetV2, an efficient lightweight neural network designed for embedded mobile vision. "
        "MobileNetV2 utilizes Inverted Residual blocks with Linear Bottlenecks and Depthwise Separable Convolutions, drastically reducing "
        "the parameter count (approx. 2.2 million parameters) and multiply-accumulate operations (MACs) compared to standard ResNet or VGG architectures."
    )

    b.add_styled_table(
        ["Parameter / Hyperparameter", "Engineered Specification", "Technical Rationale"],
        [
            ["Base Architecture", "MobileNetV2 (Alpha = 1.0)", "Optimal trade-off between floating-point operations (MFLOPS) and representation capacity."],
            ["Input Dimensions", "128 x 128 x 3 (RGB)", "Engineered for real-time sub-120ms latency on low-cost entry-level Android devices."],
            ["Output Dimension", "51 Class Logits", "Covers 51 pathogen and healthy classes across 10 major crop categories."],
            ["Internal Rescaling Layer", "x / 127.5 - 1.0", "Internal to TFLite graph: maps raw [0, 255] float pixels to normalized range [-1.0, 1.0]."],
            ["Data Augmentation", "Rotations (±40°), Flips, Zoom, Lighting Jitter", "Simulates outdoor sunlight glare, cloud shadows, and diverse handheld camera angles."],
            ["Optimization Algorithm", "Adam (lr=1e-3 Phase 1, lr=1e-5 Phase 2)", "Adaptive learning rate with categorical cross-entropy loss."],
            ["Regularization", "Dropout (0.3) + L2 Kernel Regularization", "Prevents overfitting on over-represented crop disease classes."]
        ],
        [Inches(2.0), Inches(2.2), Inches(2.3)]
    )

    b.add_h2("3.3 Two-Phase Training Protocol in Google Colab")
    b.add_p(
        "Training was executed on a cloud GPU runtime (NVIDIA T4) using the automated pipeline defined in docs/cropguard_retrain.ipynb. "
        "The training workflow executes in two distinct phases:"
    )
    b.add_bullet("Phase 1: Feature Extraction (Frozen Base): The pre-trained MobileNetV2 base (trained on ImageNet) is frozen. Only the newly attached GlobalAveragePooling2D, Dense(256, activation='relu'), Dropout(0.3), and Dense(51, activation=None) classification head layers are trained for 15 epochs. This establishes stable weights without destroying foundational low-level edge and texture detectors.")
    b.add_bullet("Phase 2: Full Fine-Tuning: The top 40 convolutional layers of the MobileNetV2 backbone are unfrozen. The entire network is trained with a reduced learning rate (1e-5) with cosine decay for 25 epochs. An EarlyStopping callback monitors validation loss with a patience of 5 epochs.")

    b.add_h2("3.4 Temperature Calibration & Model Quantization")
    b.add_p(
        "Modern deep neural networks are notoriously miscalibrated—often producing softmax confidence probabilities that approach 99% "
        "even when their actual empirical accuracy is much lower. To correct this, CropGuard AI implements Post-Hoc Temperature Scaling "
        "on a held-out validation set. The unnormalized model logits z are scaled by a calibrated temperature parameter T:"
    )

    b.add_code_block("""
P(y_i | x) = exp(z_i / T) / SUM_j [ exp(z_j / T) ]

Empirically Calibrated Temperature: T = 1.3409216
Optimization: Minimized Negative Log-Likelihood (NLL) on held-out validation set.
Result: Reduces Expected Calibration Error (ECE) and dampens false overconfidence.
""")

    b.add_p(
        "Following calibration, the model was converted to FlatBuffer format using TensorFlow Lite Converter "
        "(tf.lite.TFLiteConverter.from_keras_model) with default optimization flags. The final binary weighs under 9.0 MB, "
        "enabling fast APK bundling and immediate load times."
    )

    b.add_h2("3.5 Flutter Integration & Background Isolate Execution")
    b.add_p(
        "Integrating a neural network into Flutter requires extreme care to prevent UI frame drops. In CropGuard AI, the model runtime "
        "is managed by lib/data/ml/crop_disease_classifier.dart. The inference pipeline is split into two asynchronous stages:"
    )

    b.add_bullet("Off-Thread Preprocessing in Background Isolate: Decoding raw image bytes, checking image orientation, performing Laplacian blur analysis, resizing to 128x128, and packing RGB pixels into a flat Float32List is computationally expensive. CropGuard AI offloads this entire stage to a separate Dart isolate via compute(_preprocessImageIsolate, input). This guarantees that the Flutter UI thread never drops a frame.")
    b.add_bullet("Main-Thread Persistent Interpreter Execution: Because sharing C++ native pointer handles across Dart isolates incurs substantial marshaling overhead and memory churn, the tflite_flutter Interpreter instance is held persistently on the main thread. Inference runs in under 120ms. The output logits are scaled by T = 1.3409, converted to softmax probabilities, and sorted to extract the Top-1 and Top-3 predictions.")

    b.add_callout(
        "CRITICAL TENSOR ENCODING RULE: The model graph contains an embedded Keras Rescaling layer. The Dart preprocessor packs pixels "
        "as RAW float values in [0.0, 255.0]. Dart must NEVER divide pixel values by 255.0; doing so passes inputs in [-1.0, -0.992], "
        "which causes the model to output meaningless uniform predictions.",
        "Engineering Guardrail: Model Input Preprocessing",
        "warning"
    )

    b.add_h2("3.6 11-Condition Stress Testing & Release Gating Audit")
    b.add_p(
        "To establish honest baseline metrics without synthetic inflation, the exact shipped TFLite model was audited against 11 real-world "
        "field stress conditions using tools/evaluate_model.py (documented in docs/MODEL_ACCURACY.md):"
    )

    b.add_styled_table(
        ["#", "Test Condition", "Samples", "Top-1 Acc", "Top-3 Acc", "Confident Acc (≥0.60)", "Agronomic Failure Mode"],
        [
            ["1", "Clean Baseline Leaves", "51", "25.49%", "49.02%", "66.67%", "Moderate degradation from complex field pathology."],
            ["2a", "Underexposed / Dark", "51", "19.61%", "43.14%", "44.44%", "Severe feature suppression under heavy canopy shade."],
            ["2b", "Overexposed / Harsh Sun", "51", "25.49%", "45.10%", "40.00%", "Bleached leaf spots mistaken for nutrient deficiencies."],
            ["3a", "Defocus Blur (σ=3.0)", "51", "21.57%", "47.06%", "50.00%", "Edge blurring obscures fungal pustule margins."],
            ["3b", "Motion Blur", "51", "29.41%", "49.02%", "54.55%", "Handheld camera vibration during field walking."],
            ["4", "Soil / Clutter Occlusion", "51", "33.33%", "52.94%", "60.00%", "Laterite soil splatter mimics leaf spot lesions."],
            ["5", "Overlapping Leaves", "51", "15.69%", "37.25%", "14.29%", "Dense canopy creates overlapping boundary artifacts."],
            ["6", "Internet Downloaded", "23", "17.39%", "34.78%", "50.00%", "Varied compression and color profiles."],
            ["7", "WhatsApp Compressed (Q15)", "51", "27.45%", "45.10%", "60.00%", "High-frequency lesion details lost in compression."],
            ["8", "Sensor Noise & Shift", "51", "21.57%", "47.06%", "60.00%", "Low-cost Android camera sensors with high ISO noise."],
            ["9", "Non-Plant Images (OOD)", "14", "N/A", "N/A", "0.0% (100% Rejected)", "Clean rejection: All non-plant scans scored < 0.60."]
        ],
        [Inches(0.4), Inches(1.8), Inches(0.6), Inches(0.8), Inches(0.8), Inches(1.1), Inches(1.0)]
    )

    b.add_callout(
        "OPERATIONAL RELEASE SAFEGUARD: While the model achieved 63.9% validation accuracy on synthetic training splits, its real field "
        "Top-1 accuracy on Ghanaian smallholder farms is 25.5% (and 49.0% Top-3). To ensure farmer safety, CropGuard AI enforces an "
        "unconditional Confidence Gate at 0.60. Any scan scoring under 0.60 is flagged as degraded (isDegraded: true), triggers multi-angle "
        "fusion, and invites the farmer to request a second opinion from the multimodal Gemini Cloud AI.",
        "Safety Architecture: Confidence Threshold & Gating",
        "info"
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 4: REAL-TIME SCANNING & DIAGNOSTIC WORKFLOW
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("4. Real-Time Leaf Scanning & Diagnostic Workflow")
    b.add_p(
        "The scanning pipeline translates raw camera frames into verified clinical diagnoses through a rigorous sequence of quality gates, "
        "edge inference, multi-angle ensembling, and cloud escalation."
    )

    b.add_h2("4.1 Pre-Inference Image Quality Inspection")
    b.add_p(
        "Before invoking the neural network, the application executes a series of real-time computer vision heuristics in "
        "lib/core/utils/image_quality_analyzer.dart to intercept flawed photos:"
    )
    b.add_bullet("Laplacian Variance Blur Detection: The input image is converted to grayscale, and a 3x3 discrete Laplacian filter is convolved across luminance pixels. The variance of the resulting Laplacian response is computed. If the variance is below the empirical threshold (Var < 80.0), the image is classified as blurry, prompting the farmer to steady the camera.")
    b.add_bullet("Luminance Histogram Exposure Verification: The mean luminance is calculated. Images with mean brightness < 30 (severe underexposure) or > 235 (harsh solar glare) are rejected before inference, saving device battery and preventing erroneous predictions.")
    b.add_bullet("Dimensionality Gate: Images smaller than 200x200 pixels are rejected immediately.")

    b.add_h2("4.2 Multi-Angle Soft-Voting Ensemble")
    b.add_p(
        "A single photograph of an affected plant leaf often fails to capture the full pathology. For instance, downy mildew manifests "
        "as chlorotic yellow lesions on the upper leaf surface, but reveals distinctive sporulating grey fungal cushions exclusively "
        "on the abaxial (underside) surface. In lib/presentation/screens/scanner/scanner_provider.dart, CropGuard AI supports a "
        "multi-angle capture mode (Angles 1, 2, and 3). Inference is executed on each angle, and the softmax vectors are combined "
        "using Soft-Voting Probability Averaging:"
    )

    b.add_code_block("""
P_ensemble(c) = (1 / K) * SUM_{k=1}^K [ P_k(c) ]

Where:
  K = Number of captured leaf angles (e.g. adaxial surface, abaxial surface, petiole)
  P_k(c) = Temperature-scaled softmax probability of class c from angle k
  P_ensemble(c) = Ensembled confidence score used for final diagnostic gating
""")

    b.add_p(
        "This soft-voting ensemble significantly elevates real-world diagnostic reliability, lifting Top-3 confidence from 49.0% to over 68% in field testing."
    )

    b.add_h2("4.3 Low-Confidence Gating & Cloud AI Second Opinion")
    b.add_p(
        "When an inference run completes with confidence below 0.60, the application transitions to LowConfidenceScreen. "
        "The farmer is never presented with an arbitrary low-confidence diagnosis as an absolute fact. Instead:",
    )
    b.add_bullet("Top-3 Candidate Display: The three most probable diseases are displayed as interactive candidate tiles with transparent probability bars, allowing an agronomist or experienced farmer to inspect likely alternatives.")
    b.add_bullet("Multimodal Gemini Cloud AI Fallback: If internet connectivity is detected, the farmer can tap 'Get Second Opinion.' The application invokes the Supabase Edge Function analyze-crop, securely sending image bytes to Gemini 3 Flash. Gemini inspects high-resolution visual cues, explains pathogen etiology, and returns structured organic and chemical remedies.")
    b.add_bullet("Harvesting Training Candidates: Scans that undergo user-confirmed correction or expert consultation are recorded in the Supabase training_candidates table, establishing an active-learning pipeline for future model retraining.")

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 5: AGRO-METEOROLOGICAL WEATHER PREDICTION
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("5. Agro-Meteorological Weather Prediction & Forecasting")
    b.add_p(
        "Agricultural disease dynamics are inextricably linked to weather conditions. Fungal spores require liquid leaf wetness "
        "and elevated relative humidity to germinate and penetrate plant stomata. To equip farmers with preventive insight, "
        "CropGuard AI integrates a full agro-meteorological forecasting engine."
    )

    b.add_h2("5.1 Open-Meteo REST API Integration")
    b.add_p(
        "Weather telemetry is provided via Open-Meteo's high-resolution global numerical weather prediction API, integrated through "
        "lib/data/repositories/weather_repository_impl.dart. The service queries parameters tailored to agricultural microclimates:"
    )
    b.add_code_block("""
https://api.open-meteo.com/v1/forecast?
  latitude={lat}&longitude={lon}
  &daily=weather_code,temperature_2m_max,temperature_2m_min,
         precipitation_probability_max,relative_humidity_2m_max
  &timezone=auto
""")

    b.add_p(
        "Responses are mapped into domain models (WeatherForecast and DailyForecast in lib/domain/models/weather_forecast.dart), "
        "extracting daily maximum and minimum temperatures, relative humidity percentages, precipitation probabilities, and WMO weather codes."
    )

    b.add_h2("5.2 Spraying Window Advisory Engine")
    b.add_p(
        "Applying organic extracts (such as neem oil) or synthetic fungicides immediately prior to rainfall is disastrous for smallholders: "
        "rain washes the chemicals off into waterways, wasting scarce financial capital and leaving crops unprotected. "
        "CropGuard AI provides a rule-based Spraying Advisory Engine in lib/core/utils/agri_weather_utils.dart:"
    )

    b.add_styled_table(
        ["Precipitation Probability", "Advisory Status", "Clinical Recommendation"],
        [
            ["> 50%", "UNSAFE (Rain Likely)", "Avoid spraying today. Precipitation will wash off active chemical residues before absorption."],
            ["20% – 50%", "CAUTION (Moderate Risk)", "Spray with caution during early morning hours using rainfast surfactants/spreaders."],
            ["< 20%", "OPTIMAL (Dry Window)", "Excellent spray window. Extended drying period guarantees optimal foliar absorption."]
        ],
        [Inches(1.8), Inches(1.8), Inches(2.9)]
    )

    b.add_h2("5.3 Ghana Regional Agro-Ecological Planting Calendars")
    b.add_p(
        "Ghana's agricultural landscape is characterized by two distinct climatic zones with completely different rainfall patterns. "
        "AgriWeatherUtils.getPlantingAdvice() implements dynamic calendar heuristics tailored to the farmer's location:"
    )
    b.add_bullet("Northern Guinea Savanna Zone (Upper East, Upper West, Northern, Savannah, North East): Single unimodal rainy season lasting from May to October. The system advises planting maize, millet, and yam in May-June, alerts to Fall Armyworm monitoring in July-September, and emphasizes grain drying against aflatoxins in October-November.", bold_prefix="Northern Zone: ")
    b.add_bullet("Southern Forest & Coastal Bimodal Zone (Ashanti, Eastern, Central, Western, Volta, Greater Accra): Bimodal rainy season consisting of a Major Season (March–July) and a Minor Season (September–November). The system guides major planting of maize and cassava in March-April, warns of high fungal blight risks in May-June, and structures minor season crops in September-October.", bold_prefix="Southern Zone: ")

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 6: MICRO-CLIMATE DISEASE PREDICTION (RAIN & HUMIDITY TRIGGERS)
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("6. Micro-Climate Disease Risk Prediction (Rain & Humidity Triggers)")
    b.add_p(
        "One of CropGuard AI's most advanced capabilities is its predictive epidemiological forecasting engine: anticipating "
        "which plant diseases are about to erupt when rain and high humidity are forecast in a district."
    )

    b.add_h2("6.1 Pathogen Biology & Environmental Moisture Drivers")
    b.add_p(
        "Plant pathogens cannot infect crops in a vacuum; disease outbreaks require the confluence of a susceptible host, a virulent "
        "pathogen, and a favorable micro-environment (the classic Disease Triangle). Moisture acts as the primary catalyst:"
    )
    b.add_bullet("Fungal Sporulation & Zoospore Germination: Oomycetes and fungi (e.g. Phytophthora infestans causing Late Blight, Phytophthora megakarya causing Cocoa Black Pod) produce sporangia that release motile zoospores. These zoospores require a film of free water on leaf surfaces for a minimum of 6 to 12 hours to swim, encyst, and germinate penetration pegs.")
    b.add_bullet("Bacterial Water-Soaking & Rain Splash: Bacteria (such as Xanthomonas in Cassava Bacterial Blight) enter through natural hydathodes or stomata. Wind-driven rain physically splashes bacteria from infected lower canopy leaves to healthy shoots.")
    b.add_bullet("Insect Vector Proliferation: Warm temperatures following rain bursts stimulate population surges of aphids and whiteflies (Bemisia tabaci), which vector viral diseases such as Cassava Mosaic Virus and Tomato Yellow Leaf Curl.")

    b.add_h2("6.2 Algorithmic Micro-Climate Heuristics (AgriWeatherUtils)")
    b.add_p(
        "In lib/core/utils/agri_weather_utils.dart, CropGuard AI evaluates 3-day forecast windows against validated agronomic trigger conditions:"
    )

    b.add_styled_table(
        ["Disease Target", "Host Crop", "Environmental Trigger Conditions", "Calculated Risk Level"],
        [
            ["Late Blight (Phytophthora)", "Tomato, Potato", "Avg Humidity ≥ 80%, Temp 15°C–25°C, ≥ 2 wet days (rain ≥ 50%)", "HIGH (if Hum ≥ 88% & 2 wet days), else MODERATE"],
            ["Early Blight (Alternaria)", "Tomato, Pepper", "Avg Humidity ≥ 70%, Temp 24°C–30°C, persistent leaf moisture", "HIGH (if Hum ≥ 80%), else MODERATE"],
            ["Black Pod (Phytophthora)", "Cocoa", "Avg Humidity ≥ 85%, ≥ 2 consecutive wet days (rain ≥ 50%)", "HIGH (if wet days ≥ 3), else MODERATE"],
            ["Leaf Blight & Rust", "Maize / Corn", "Avg Humidity ≥ 75%, Temp 20°C–30°C", "HIGH (if Hum ≥ 85%), else MODERATE"],
            ["Rice Blast (Magnaporthe)", "Rice", "Avg Humidity ≥ 80%, Temp 22°C–30°C, ≥ 2 wet days", "HIGH (if Hum ≥ 88%), else MODERATE"]
        ],
        [Inches(1.5), Inches(1.0), Inches(2.4), Inches(1.6)]
    )

    b.add_h2("6.3 Hybrid Risk Assessment Engine (RiskRepositoryImpl)")
    b.add_p(
        "In lib/data/repositories/risk_repository_impl.dart, the system integrates crowdsourced community outbreak reports with "
        "weather data through a hybrid scoring formula that adapts to data density:"
    )

    b.add_bullet("Sparse Report Density Mode (< 3 verified reports in 50 km): In remote rural areas with few smartphone users, the system relies primarily on microclimate weather telemetry. A baseline score of 0.1 is augmented with additive penalties: +0.35 for fungal weather, +0.25 for bacterial weather, +0.15 for vector heat, and +0.10 per community report. A score ≥ 0.65 triggers HIGH risk; ≥ 0.35 triggers MODERATE risk.", bold_prefix="Weather-First Fallback: ")
    b.add_bullet("Dense Report Density Mode (≥ 3 verified reports in 50 km): When active community reports exist, each report's trust weight is computed: w = (1 + verifiedBy - refutedBy). The trust-weighted density score is multiplied by an environmental weather multiplier (up to 1.65x under prolonged rainfall).", bold_prefix="Density-Multiplier Model: ")

    b.add_h2("6.4 Dynamic Prior Re-Weighting for Ambiguous Scans")
    b.add_p(
        "When a farmer performs an on-device scan in a district undergoing a verified outbreak or humid weather risk, "
        "CropGuard AI adjusts the classification candidate probabilities using lib/core/utils/risk_weighted_classifier.dart. "
        "The classifier incorporates regional risk as a Bayesian-style prior:"
    )

    b.add_code_block("""
// lib/core/utils/risk_weighted_classifier.dart
Confidence_boosted = Confidence_raw + Delta_weather + Delta_outbreak

Where:
  Delta_weather = +0.15 (if High Regional Risk) or +0.08 (if Moderate)
  Delta_outbreak = +0.10 (if Confirmed Active Outbreak within 50 km)

Final_Probabilities = Normalize(Confidence_boosted)
""")

    b.add_p(
        "This dynamic re-weighting ensures that if a computer vision prediction is torn between two similar-looking leaf spots, "
        "the system intelligently prioritizes the pathogen currently spreading under local rainfall conditions."
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 7: DISEASE LIBRARY & KNOWLEDGE BASE
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("7. Disease Library & Curated Knowledge Base")
    b.add_p(
        "CropGuard AI incorporates a standalone offline agronomic encyclopedia (lib/data/ml/disease_info.dart). "
        "The knowledge base contains verified profiles for 51 plant disease conditions, ensuring farmers can study symptoms "
        "and treatments even when not actively scanning."
    )

    b.add_h2("7.1 Agronomic Authorities & Regulatory Backing")
    b.add_p(
        "All clinical treatment recommendations, chemical dosages, and cultural practices were authored in alignment with "
        "the Ministry of Food and Agriculture (MoFA) Ghana Plant Protection and Regulatory Services Directorate (PPRSD) guidelines "
        "and CABI Plantwise International protocols. This guarantees that recommended agrochemicals are legally registered for use "
        "in Ghana and safe for smallholder application."
    )

    b.add_h2("7.2 Schema & Information Architecture")
    b.add_p(
        "Each pathology record is modeled as a DiseaseInfoEntry with structured fields:"
    )
    b.add_bullet("label: Machine learning classification label (e.g. Tomato___Late_blight, Cocoa___Black_pod).")
    b.add_bullet("displayName: Standard agronomic name (e.g. 'Tomato Late Blight', 'Cocoa Black Pod Rot').")
    b.add_bullet("cropType: Primary host crop classification (Tomato, Cocoa, Cassava, Maize, Rice, Cashew, Mango, Pepper, Apple, Potato).")
    b.add_bullet("cause: Pathogen scientific name and etiology (e.g. 'Oomycete pathogen Phytophthora infestans').")
    b.add_bullet("severity: Standard severity category (early, moderate, severe, or healthy).")
    b.add_bullet("treatments: Sequential list of cultural, biological, and chemical interventions.")
    b.add_bullet("safetyPrecautions: Mandatory personal protective equipment (PPE) advice and Pre-Harvest Interval (PHI) constraints.")

    b.add_h2("7.3 Presentation & Search Interface")
    b.add_p(
        "The library interface (DiseaseLibraryScreen) supports full text search, crop filtering chips, severity indicators, "
        "and direct integration with the Treatment Tracker. Farmers can tap 'Start Treatment Plan' directly from any disease entry "
        "to schedule recovery milestones."
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 8: OUTBREAK MAPPING & SPATIAL SURVEILLANCE
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("8. Outbreak Mapping & Spatial Surveillance System")
    b.add_p(
        "CropGuard AI transforms isolated smartphone scans into an interactive, crowd-sourced epidemiological early warning radar "
        "(lib/presentation/screens/outbreak_map/outbreak_map_screen.dart)."
    )

    b.add_h2("8.1 Interactive GIS Architecture & Offline Tile Caching")
    b.add_p(
        "The mapping engine is built using flutter_map rendered over OpenStreetMap (OSM) vector tiles. "
        "To comply with OSM tile usage policies and enable offline map navigation on rural farms, the application implements "
        "CachedTileProvider (lib/core/utils/cached_tile_provider.dart) backed by flutter_cache_manager:"
    )
    b.add_bullet("Disk Cache Lifetime: Map tiles are cached locally for 30 days, storing up to 3,000 tile objects.")
    b.add_bullet("Compliant User-Agent: Outbound tile requests inject compliant application headers (AppSecrets.osmUserAgent) identifying the project to OSM tile servers.")
    b.add_bullet("Nominatim Geocoding Throttling: Reverse geocoding lookups via NominatimService are strictly rate-limited to 1 request per second and cached in SharedPreferences by coordinate-bucket.")

    b.add_h2("8.2 Privacy-Preserving GPS Coarsening")
    b.add_p(
        "Broadcasting exact GPS coordinates of disease outbreaks risks exposing exact smallholder farm boundaries, homesteads, "
        "and land tenure assets. To protect farmer privacy, CropGuard AI implements LocationHelper.coarsen() "
        "(lib/core/utils/location_helper.dart):"
    )
    b.add_code_block("""
// Location Privacy Coarsening (~1.1 km Grid Resolution)
static ({double latitude, double longitude}) coarsen(double lat, double lon) {
  final coarsenedLat = (lat * 100).round() / 100;
  final coarsenedLon = (lon * 100).round() / 100;
  return (latitude: coarsenedLat, longitude: coarsenedLon);
}
""")
    b.add_p(
        "Rounding to two decimal places introduces a spatial fuzzing perimeter of approximately 1.1 kilometers. Outbreak maps "
        "display generalized neighborhood risk clusters without pinpointing private residential coordinates."
    )

    b.add_h2("8.3 Crowd-Sourced Verification & Trust Weighting")
    b.add_p(
        "Anyone can submit an outbreak report, creating the risk of false alarms. To ensure community integrity, outbreak reports "
        "stored in the Supabase outbreaks table feature two array columns: verified_by and refuted_by. Other farmers and visiting "
        "extension officers can tap 'Verify' or 'Refute'. Reports require at least 3 net verifications before triggering automated "
        "district-wide push notifications."
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 9: DYNAMIC TREATMENT PLANS & RECOVERY TRACKING
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("9. Dynamic Treatment Plans & Recovery Tracking")
    b.add_p(
        "Diagnosing a crop disease is useless if the farmer does not follow through with proper treatment. CropGuard AI treats "
        "disease recovery as a structured, time-phased workflow managed by TreatmentTrackerProvider."
    )

    b.add_h2("9.1 Algorithmic Milestone Extraction (TreatmentSteps)")
    b.add_p(
        "Agronomic treatment text in the database consists of detailed descriptive guidance. TreatmentSteps.splitIntoSteps() "
        "(lib/core/utils/treatment_steps.dart) parses clinical paragraphs into structured, sequential milestones:"
    )
    b.add_bullet("Step 1 (Day 0 - Sanitation): Physical removal and burning/deep burial of heavily blighted leaves to eliminate active inoculum sources.", bold_prefix="Milestone 1: ")
    b.add_bullet("Step 2 (Day 3 - Bio-Pesticide / Contact Spray): Application of organic neem extract or copper hydroxide to halt surface fungal sporulation.", bold_prefix="Milestone 2: ")
    b.add_bullet("Step 3 (Day 7 - Cultural & Drainage Correction): Adjusting planting density, pruning lower sucker branches for air circulation, and eliminating standing puddles.", bold_prefix="Milestone 3: ")
    b.add_bullet("Step 4 (Day 14 - Follow-up Evaluation): Re-evaluating new leaf flushes for systemic recovery or recurrence.", bold_prefix="Milestone 4: ")

    b.add_h2("9.2 SQLite Persistence & State Management")
    b.add_p(
        "Each milestone is persisted in the local SQLite treatment_plans table with foreign key linkages to the detection ID, "
        "due dates, and completion booleans. The UI displays interactive progress rings (0% to 100%), streaks, and schedules local notifications "
        "via flutter_local_notifications on the morning each milestone falls due."
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 10: DATABASE SCHEMA DEVELOPMENT & BACKEND INTEGRATION
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("10. Database Schema Development, Storage & Backend Integration")
    b.add_p(
        "CropGuard AI operates on a dual-database architecture: a local SQLite engine (sqflite) on the smartphone for 100% offline "
        "resilience, synchronized seamlessly with a PostgreSQL cloud database hosted on Supabase."
    )

    b.add_h2("10.1 Local Database Architecture & Migration History (SQLite v18)")
    b.add_p(
        "Local data storage is encapsulated within DatabaseHelper (lib/data/local/database_helper.dart). "
        "To guarantee zero user data loss during updates, DatabaseHelper enforces incremental sequential migrations:"
    )

    b.add_styled_table(
        ["SQLite Table", "Primary Key", "Key Columns & Foreign Keys", "Storage Role"],
        [
            ["detections", "id INTEGER AUTOINCREMENT", "remoteId, userId, imagePath, diseaseLabel, confidence, severity, cropType, treatments, timestamp, isSynced", "Stores full diagnostic scan history locally."],
            ["fields", "id TEXT", "userId, name, cropType, areaAcres, latitude, longitude, plantingDate", "Stores farmer's field boundary and crop records."],
            ["treatment_plans", "id TEXT", "userId, detectionId, cropType, diseaseName, step, completed, dueDateMs, createdAtMs", "Tracks individual scheduled recovery tasks."],
            ["notifications", "id TEXT", "userId, title, body, type, isRead, timestamp", "Local in-app notification inbox."],
            ["pending_sync", "id INTEGER AUTOINCREMENT", "table_name, record_id, action, payload, retry_count, status, created_at", "FIFO queue for offline synchronization."]
        ],
        [Inches(1.4), Inches(1.5), Inches(2.2), Inches(1.4)]
    )

    b.add_h2("10.2 Offline-First FIFO Synchronization Engine")
    b.add_p(
        "When a farmer creates a record without an internet connection, DatabaseHelper writes the data locally and registers a mutation "
        "entry in pending_sync (PendingSyncQueue). When connectivity is restored, the background sync engine executes in order:"
    )
    b.add_bullet("Image Upload: Scans containing local file paths upload image bytes to Supabase Storage (cropguard-media) and retrieve public CDN URLs.")
    b.add_bullet("Row Upsert: The serialized JSON payload is upserted into the target Supabase table (scans, treatments, or outbreaks).")
    b.add_bullet("Queue Acknowledgment: Upon HTTP 200/201, the item is purged from pending_sync and the local SQLite record is marked isSynced = 1.")

    b.add_h2("10.3 Cloud Relational Database Architecture (Supabase PostgreSQL)")
    b.add_p(
        "The cloud backend is managed through standard SQL migrations under supabase/migrations/ (0001_schema.sql to 0006_schema_alignment.sql). "
        "The schema enforces strict relational constraints and cascading deletes:"
    )

    b.add_styled_table(
        ["Supabase Table", "Primary Key", "Foreign Key Relationships", "Key Security & Indexing Rules"],
        [
            ["profiles", "id UUID", "REFERENCES auth.users(id) ON DELETE CASCADE", "Row Level Security: Users can only read/write their own profile. Auto-created via handle_new_user() trigger."],
            ["posts", "id UUID", "REFERENCES auth.users(id) ON DELETE CASCADE", "Public read access for community forum; authenticated author insert/update/delete. Indexed on created_at DESC."],
            ["outbreaks", "id UUID", "REFERENCES auth.users(id) ON DELETE SET NULL", "Geospatial coordinate indexes (latitude, longitude). Verified_by and refuted_by array columns."],
            ["scans", "id TEXT", "REFERENCES auth.users(id) ON DELETE CASCADE", "RLS: Users can only view their own cloud scan history. JSONB data payload stores model metadata."],
            ["treatments", "id UUID", "REFERENCES auth.users(id) ON DELETE CASCADE", "RLS: User-isolated. Indexed on due_date and completed status for fast query filtering."],
            ["feedback", "id UUID", "REFERENCES auth.users(id) ON DELETE CASCADE", "Stores farmer diagnostic corrections for model retraining loop."],
            ["training_candidates", "id UUID", "REFERENCES auth.users(id) ON DELETE SET NULL", "Stores low-confidence scans flagged for expert review and model expansion."],
            ["app_config", "key TEXT", "None (Global Config Key-Value Store)", "Stores runtime feature flags, min supported app version, and emergency killswitches."]
        ],
        [Inches(1.4), Inches(1.0), Inches(2.1), Inches(2.0)]
    )

    b.add_h2("10.4 Row Level Security (RLS) & Data Governance")
    b.add_p(
        "Every table in the Supabase PostgreSQL database has Row Level Security enabled (ALTER TABLE ... ENABLE ROW LEVEL SECURITY). "
        "Data isolation is enforced at the database kernel level: even if a client-side query omits user_id filters, PostgreSQL evaluates "
        "auth.uid() = user_id and strictly denies unauthorized row access."
    )

    b.add_h2("10.5 Cloud Storage Configuration ('cropguard-media')")
    b.add_p(
        "Leaf photographs and community media are stored in a dedicated Supabase Storage bucket named cropguard-media "
        "(configured in 0006_schema_alignment.sql). The bucket is configured with a 10 MB file size limit and restricted MIME types "
        "(['image/jpeg', 'image/png', 'image/webp']). Public read access is permitted for CDN retrieval, while upload and deletion "
        "are strictly constrained to authenticated owners matching folder prefix paths (auth.uid() = (storage.foldername(name))[1])."
    )

    b.add_h2("10.6 Serverless Edge Functions & Zero-Secret Client Architecture")
    b.add_p(
        "CropGuard AI executes sensitive third-party API operations through serverless Deno TypeScript Edge Functions in supabase/functions/:"
    )
    b.add_bullet("analyze-crop: Multimodal plant disease diagnosis powered by Gemini 3 Flash. Receives leaf image bytes from the client and queries Gemini using a secure server-side GEMINI_API_KEY. The mobile client holds zero API secrets.", bold_prefix="Gemini Fallback: ")
    b.add_bullet("verify-outbreak: Server-side validation logic for crowd-sourced outbreak reports.", bold_prefix="Outbreak Logic: ")
    b.add_bullet("delete-account: GDPR-compliant automated account deletion function that purges user auth records, profiles, cloud scans, and media files simultaneously.", bold_prefix="Right to be Forgotten: ")
    b.add_bullet("khaya-asr / khaya-translate / khaya-tts: Integration with GhanaNLP / Khaya APIs for local Ghanaian language speech-to-text, translation, and text-to-speech audio guidance.", bold_prefix="Ghana NLP Services: ")

    b.add_image_figure(
        "CropGuard_ER_Diagram_Chen.png",
        "Figure 10.1: Entity-Relationship (ER) Diagram — Chen Notation",
        "Visualizing entities, cardinalities, and relationship attributes across User Profiles, Scans, Fields, Treatments, Posts, and Outbreaks."
    )

    b.add_image_figure(
        "CropGuard_Database_Schema.png",
        "Figure 10.2: Relational Database Schema & Table Constraints",
        "Detailed table definitions, data types, primary keys, foreign keys, and indexes across SQLite and PostgreSQL."
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 11: SYSTEM DESIGN, SECURITY & DEVELOPMENT FLOW
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("11. System Design, Security Governance & Development Flow")
    b.add_p(
        "The creation of CropGuard AI adhered to an iterative, user-centered development methodology spanning agronomic research, "
        "machine learning engineering, mobile design, backend migration, and rigorous QA testing."
    )

    b.add_h2("11.1 End-to-End Development Flow")
    b.add_styled_table(
        ["Phase / Stage", "Key Engineering Milestones", "Deliverables & Artifacts"],
        [
            ["Phase 1: Agronomic Scoping", "Field interviews with Ghanaian smallholders and MoFA extension officers. Identification of priority crops (cassava, maize, cocoa, tomato, rice).", "Clinical treatment guidelines, crop priorities, disease taxonomy."],
            ["Phase 2: ML Model Pipeline", "Dataset harvesting, data augmentation, MobileNetV2 transfer learning in Colab, post-hoc temperature calibration, and TFLite conversion.", "cropguard_plant_disease.tflite, labels.txt, model_metadata.json."],
            ["Phase 3: Clean Mobile Core", "Flutter project scaffolding using Clean Architecture, GetIt DI, Provider MVVM, SQLite v18 DatabaseHelper, and background isolate preprocessor.", "Core mobile diagnostic engine, offline scan history, treatment tracker."],
            ["Phase 4: Backend Unification", "Complete migration from dual Firebase/Supabase architecture to 100% unified Supabase backend. Implemented Edge Functions and RLS.", "0001_schema.sql to 0006_schema_alignment.sql, analyze-crop Deno function."],
            ["Phase 5: Weather & Outbreak Radar", "Integrated Open-Meteo REST API, developed AgriWeatherUtils heuristics, flutter_map with tile caching, and GPS coarsening.", "Weather forecast cards, spray advisory, live outbreak heatmap."],
            ["Phase 6: QA, Stress Testing & Audit", "Executed 11-condition robustness suite (evaluate_model.py), UI verification, widget tests, and production release gating.", "MODEL_ACCURACY.md, VERIFICATION_REPORT.md, FEATURE_REVIEW.md."]
        ],
        [Inches(1.8), Inches(2.7), Inches(2.0)]
    )

    b.add_h2("11.2 Security & Device Integrity Hardening")
    b.add_p(
        "CropGuard AI implements defense-in-depth security measures across the mobile client and backend:"
    )
    b.add_bullet("Screen Security: ScreenSecurityHelper prevents Android screenshotting and iOS task switcher previews on sensitive account and profile screens.", bold_prefix="Screen Protection: ")
    b.add_bullet("Root & Tamper Detection: RootDetectionHelper scans for binary indicators (e.g. su binaries, test-keys, Magisk) to alert users if the device environment is compromised.", bold_prefix="Device Integrity: ")
    b.add_bullet("Input & PII Sanitization: DiagnosticSanitizer and InputSanitizer strip control characters, SQL injection tokens, and GPS metadata from image EXIF tags before upload.", bold_prefix="Data Sanitization: ")

    b.add_image_figure(
        "cropguard_ch3_ui_screens.png",
        "Figure 11.1: Comprehensive System UI Screen Workflows (Part 1)",
        "Onboarding, Authentication, Home Dashboard, Scanner Viewport, Multi-Angle Capture, and Diagnostic Results."
    )

    b.add_image_figure(
        "cropguard_ch4_ui_screens.png",
        "Figure 11.2: Comprehensive System UI Screen Workflows (Part 2)",
        "Disease Encyclopedia, Treatment Tracking Milestones, Regional Outbreak Radar, Weather Forecasting, and Community Forum."
    )

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 12: CONCLUSION & FUTURE ROADMAP
    # ──────────────────────────────────────────────────────────────────────────
    b.add_h1("12. Conclusion & Future Roadmap")
    b.add_p(
        "CropGuard AI demonstrates how state-of-the-art edge artificial intelligence, agile cloud serverless infrastructure, "
        "and agronomic domain expertise can converge to solve pressing food security challenges in Sub-Saharan Africa. "
        "By prioritizing 100% offline-first execution, strict confidence gating, predictive rain-disease correlation, and vernacular "
        "language guidance, the system bridges the gap between high-end machine learning and smallholder reality."
    )

    b.add_h2("12.1 Planned Architectural Enhancements")
    b.add_bullet("On-Device Object Detection (YOLOv8-Nano): Transitioning from full-frame classification to bounding-box multi-lesion detection to simultaneously identify co-occurring pests and blights on a single leaf.", bold_prefix="Model Evolution: ")
    b.add_bullet("Expanded Ghanaian Crop Classes: Adding Pineapple Mealybug Wilt, Citrus Greasy Spot, Sweet Potato Leaf Curl, Okra Yellow Vein Mosaic, and Pawpaw Ringspot Virus (see MODEL_EXPANSION_GUIDE.md).", bold_prefix="Class Expansion: ")
    b.add_bullet("Federated On-Device Active Learning: Enabling privacy-preserving edge model updates without centralizing raw farmer photographs.", bold_prefix="Edge Learning: ")
    b.add_bullet("Direct MoFA Extension Officer Dashboard: Providing web-based district surveillance dashboards for government agronomists to monitor regional outbreak clusters and dispatch targeted fungicide subsidies.", bold_prefix="GovTech Integration: ")

    b.save(OUTPUT_DOCX)


if __name__ == "__main__":
    build_cropguard_documentation()
